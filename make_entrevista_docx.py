from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# Margins
from docx.shared import Cm
section = doc.sections[0]
section.left_margin = Cm(3)
section.right_margin = Cm(2)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)

# Styles
style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(11)

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    return p

def add_meta(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p

def add_block_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
    # Shading
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F0E8')
    pPr.append(shd)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    return p

def add_question(doc, num, text, notes=None):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.left_indent = Inches(0.3)
    run = p.add_run(text)
    run.font.size = Pt(11)
    if notes:
        p2 = doc.add_paragraph()
        p2.paragraph_format.left_indent = Inches(0.6)
        r = p2.add_run(f'   → {notes}')
        r.italic = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return p

def add_notes_line(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    return p

# Header
add_title(doc, 'Roteiro de Entrevista Semiestruturada')
doc.add_paragraph()

meta_lines = [
    ('Entidade: ', 'CPF Pia do Sul - Santa Maria/RS (13ª RT)'),
    ('Objetivo: ', 'Levantamento de requisitos e validação do MVP do ecossistema digital'),
    ('Duração estimada: ', '40-50 minutos'),
    ('Formato: ', 'Semiestruturada — as perguntas são um guia, não um questionário rígido'),
]
for label, value in meta_lines:
    p = doc.add_paragraph()
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(value)
    r2.font.size = Pt(11)

doc.add_paragraph()

# Bloco 1
add_block_heading(doc, 'Bloco 1 — Perfil e contexto  (~5 min)')
add_question(doc, 1, 'Quantos associados a entidade possui atualmente? Há distinção entre categorias (sócio pleno, dependente, invernada)?')
add_question(doc, 2, 'Quantas pessoas integram a diretoria/administração? Quem lida com a parte administrativa no dia a dia?')
add_question(doc, 3, 'Qual o volume médio de eventos por ano? (bailes, fandangos, festividades internas)')
add_question(doc, 4, 'Quantas pessoas comparecem em média a um evento? Há variação entre tipos de evento?')
add_question(doc, 5, 'A entidade utiliza algum sistema ou aplicativo atualmente para qualquer parte da gestão? (financeiro, cadastro, comunicação)')

# Bloco 2
add_block_heading(doc, 'Bloco 2 — Gestão de associados hoje  (~10 min)')
add_question(doc, 1, 'Como é feito hoje o cadastro de um novo associado? Existe algum formulário ou ficha?')
add_question(doc, 2, 'Onde essas informações ficam armazenadas? (planilha, caderno, sistema?)')
add_question(doc, 3, 'Como vocês sabem quais associados estão ativos, inativos ou em dia com as mensalidades?')
add_question(doc, 4, 'O que acontece quando um associado precisa consultar sua situação? Quem responde, como?')
add_question(doc, 5, 'Qual a maior dificuldade nesse processo hoje?')

# Bloco 3
add_block_heading(doc, 'Bloco 3 — Controle de mensalidades  (~10 min)')
add_question(doc, 1, 'Como é feito o registro de pagamento de mensalidade atualmente?')
add_question(doc, 2, 'Como é identificado quem está inadimplente? Com que frequência isso é verificado?')
add_question(doc, 3, 'Existe algum tipo de cobrança ou notificação para inadimplentes? Como funciona?')
add_question(doc, 4, 'Já houve problemas de informação perdida ou divergente sobre pagamentos?')

# Bloco 4
add_block_heading(doc, 'Bloco 4 — Gestão de eventos, bailes e fandangos  (~10 min)')
add_question(doc, 1, 'Como funciona a organização de um baile ou fandango do início ao fim? Quem faz o quê?')
add_question(doc, 2, 'Como é feita a reserva de mesas? (presencial, WhatsApp, lista física?)')
add_question(doc, 3, 'Como são vendidos/controlados os ingressos? Há ingresso antecipado?')
add_question(doc, 4, 'Como a entidade sabe quantas mesas/ingressos ainda estão disponíveis em tempo real?')
add_question(doc, 5, 'Como é feito o controle de quem pagou o ingresso na entrada do evento? (lista impressa, carimbo, outro?)')
add_question(doc, 6, 'Os pagamentos são feitos presencialmente, via PIX, transferência? Como é registrado que o pagamento foi recebido?')
add_question(doc, 7, 'Já houve situações de conflito de reserva, mesa dupla vendida ou ingresso sem controle?')
add_question(doc, 8, 'Como é feita a prestação de contas financeira após um evento?')

# Bloco 5
add_block_heading(doc, 'Bloco 5 — Perfil dos usuários do sistema  (~5 min)')
add_question(doc, 1, 'Quem você imagina usando o painel administrativo no dia a dia? Essa pessoa tem familiaridade com sistemas digitais?')
add_question(doc, 2, 'E os associados — qual faixa etária predomina? Eles usam smartphone com facilidade?')
add_question(doc, 3, 'Há associados que precisariam de ajuda para usar um aplicativo? Como a entidade lida com isso hoje?')

# Bloco 6
add_block_heading(doc, 'Bloco 6 — Expectativas e prioridades  (~10 min)')
add_question(doc, 1, 'Se você pudesse resolver só um problema com esse sistema, qual seria?')
add_question(doc, 2, 'Das funcionalidades previstas (associados, mensalidades, eventos, reservas, app do associado) — qual você considera mais urgente?')
add_question(doc, 3, 'Tem alguma coisa que não mencionamos mas que seria importante ter?')
add_question(doc, 4, 'Existe alguma preocupação sua em relação à adoção do sistema pela equipe ou pelos associados?')
add_question(doc, 5, 'O que você esperaria ver funcionando no sistema para considerar que foi um sucesso?')

# Observacoes
doc.add_paragraph()
add_block_heading(doc, 'Observações para condução')
notes = [
    'Solicitar autorização para gravação de áudio antes de iniciar',
    'Levar impresso o diagrama de arquitetura ou rascunho de telas para provocar feedback visual',
    'O Bloco 6 é o mais estratégico para definição do MVP — não encurtar',
    'Anotar exemplos concretos citados (ex.: "isso aconteceu no baile de agosto") — são dados qualitativos valiosos',
    'Se surgirem processos não mapeados, explorar com: "Como funciona exatamente? Quem faz? Com qual frequência?"',
]
for note in notes:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.3)
    r = p.add_run(note)
    r.font.size = Pt(11)

# Space for notes
doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run('Anotações durante a entrevista:')
r.bold = True
r.font.size = Pt(11)
for _ in range(20):
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(14)
    # Underline via border bottom
    pPr = p2._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'AAAAAA')
    pBdr.append(bottom)
    pPr.append(pBdr)

doc.save(r'e:\TCC\docs\entrevista_cpf_pia_do_sul.docx')
print('Word gerado com sucesso')
