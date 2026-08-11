from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

section = doc.sections[0]
section.left_margin = Cm(3)
section.right_margin = Cm(2)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(11)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    return p


def add_meta(doc, label, value):
    p = doc.add_paragraph()
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(value)
    r2.font.size = Pt(11)
    return p


def add_quote(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    return p


def add_block_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F0E8')
    pPr.append(shd)
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_question(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.bold = True
    return p


def add_answer_lines(doc, n=4):
    for _ in range(n):
        ans = doc.add_paragraph()
        ans.paragraph_format.left_indent = Inches(0.3)
        ans.paragraph_format.space_after = Pt(10)
        pPr = ans._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'AAAAAA')
        pBdr.append(bottom)
        pPr.append(pBdr)
        ans.add_run(' ')


# Header
add_title(doc, 'Segunda Rodada — Sobre os Sistemas Já Utilizados')
doc.add_paragraph()

add_meta(doc, 'Entidade: ', 'CPF Pia do Sul - Santa Maria/RS (13ª RT)')
add_meta(doc, 'Objetivo: ', 'Esclarecer três respostas do primeiro formulário que mencionaram "sistema", para entender se é a mesma ferramenta ou ferramentas diferentes, e o que cada uma cobre')
add_meta(doc, 'Formato: ', 'Respostas por escrito, direto abaixo de cada pergunta')

doc.add_paragraph()

points = [
    (
        'Sobre o "sistema em teste" (Bloco 1)',
        'Pergunta original: "A entidade utiliza algum sistema ou aplicativo atualmente para qualquer parte da gestão?"\nResposta de vocês: "Há um Sistema em teste."',
        'Esse sistema tem nome? Foi feito sob encomenda pra vocês ou é uma ferramenta pronta de mercado? Ele é pago? E hoje ele está realmente em uso no dia a dia, ou ainda em fase de avaliação/teste mesmo?',
    ),
    (
        'Sobre o "sistema eletrônico" do cadastro (Bloco 2)',
        'Pergunta original: "Como é feito hoje o cadastro de um novo associado?"\nResposta de vocês: "Sistema de ficha que é transferida para o Sistema eletrônico."',
        'Esse é o mesmo sistema mencionado na pergunta anterior, ou é outro? Quem tem acesso a ele — só a secretaria, ou os associados também conseguem consultar seus próprios dados nele? Ele permite exportar os dados dos associados (planilha, CSV etc.), caso um dia precisem migrar de sistema?',
    ),
    (
        'Sobre o "sistema informatizado" da mensalidade (Bloco 3)',
        'Pergunta original: "Como é feito o registro de pagamento de mensalidade atualmente?"\nResposta de vocês: "Controle via conta bancária e Sistema informatizado."',
        'Esse sistema é o mesmo dos dois pontos anteriores, ou é uma ferramenta separada (por exemplo, algo do próprio banco)? É pago? Se o sistema deste projeto ficar pronto e cobrir isso melhor, vocês teriam interesse em substituir esse sistema atual, ou prefeririam que os dois continuassem coexistindo, cada um cuidando de uma parte?',
    ),
]

for heading, quote, question in points:
    add_block_heading(doc, heading)
    add_quote(doc, quote)
    add_question(doc, question)
    add_answer_lines(doc, 4)

doc.save(r'TCC-FINAL\respostas_formulario_inicial\contraperguntas_sistemas.docx')
print('Word gerado com sucesso')
