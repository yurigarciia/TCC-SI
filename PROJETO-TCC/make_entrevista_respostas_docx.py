from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

section = doc.sections[0]
section.left_margin = Cm(3)
section.right_margin = Cm(2)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Arial'
style.font.size = Pt(11)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    return p


def add_block_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F0E8')
    pPr.append(shd)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_question_with_answer_space(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.bold = True

    label = doc.add_paragraph()
    label.paragraph_format.left_indent = Inches(0.3)
    label.paragraph_format.space_after = Pt(2)
    r = label.add_run('Resposta:')
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Three blank ruled lines under each question for a typed answer.
    for _ in range(3):
        ans = doc.add_paragraph()
        ans.paragraph_format.left_indent = Inches(0.3)
        ans.paragraph_format.space_after = Pt(10)
        pPr = ans._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '4')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'AAAAAA')
        pBdr.append(bottom)
        pPr.append(pBdr)
        ans.add_run(' ')
    return p


# Header
add_title(doc, 'Roteiro de Entrevista — Respostas Escritas')
doc.add_paragraph()

meta_lines = [
    ('Entidade: ', 'CPF Pia do Sul - Santa Maria/RS (13ª RT)'),
    ('Objetivo: ', 'Levantamento de requisitos e validação do MVP do ecossistema digital'),
    ('Formato: ', 'Respostas por escrito — pode responder direto neste documento, abaixo de cada pergunta'),
    ('Instruções: ', 'Não precisa responder tudo em detalhe; frases curtas já ajudam. Se uma pergunta não fizer sentido pra realidade de vocês, pode pular ou comentar o motivo.'),
]
for label, value in meta_lines:
    p = doc.add_paragraph()
    r1 = p.add_run(label)
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(value)
    r2.font.size = Pt(11)

doc.add_paragraph()

blocks = [
    ('Bloco 1 — Perfil e contexto', [
        'Quantos associados a entidade possui atualmente? Há distinção entre categorias (sócio pleno, dependente, invernada)?',
        'Quantas pessoas integram a diretoria/administração? Quem lida com a parte administrativa no dia a dia?',
        'Qual o volume médio de eventos por ano? (bailes, fandangos, festividades internas)',
        'Quantas pessoas comparecem em média a um evento? Há variação entre tipos de evento?',
        'A entidade utiliza algum sistema ou aplicativo atualmente para qualquer parte da gestão? (financeiro, cadastro, comunicação)',
    ]),
    ('Bloco 2 — Gestão de associados hoje', [
        'Como é feito hoje o cadastro de um novo associado? Existe algum formulário ou ficha?',
        'Onde essas informações ficam armazenadas? (planilha, caderno, sistema?)',
        'Como vocês sabem quais associados estão ativos, inativos ou em dia com as mensalidades?',
        'O que acontece quando um associado precisa consultar sua situação? Quem responde, como?',
        'Qual a maior dificuldade nesse processo hoje?',
    ]),
    ('Bloco 3 — Controle de mensalidades', [
        'Como é feito o registro de pagamento de mensalidade atualmente?',
        'Como é identificado quem está inadimplente? Com que frequência isso é verificado?',
        'Existe algum tipo de cobrança ou notificação para inadimplentes? Como funciona?',
        'Já houve problemas de informação perdida ou divergente sobre pagamentos?',
    ]),
    ('Bloco 4 — Gestão de eventos, bailes e fandangos', [
        'Como funciona a organização de um baile ou fandango do início ao fim? Quem faz o quê?',
        'Como é feita a reserva de mesas? (presencial, WhatsApp, lista física?)',
        'Como são vendidos/controlados os ingressos? Há ingresso antecipado?',
        'Como a entidade sabe quantas mesas/ingressos ainda estão disponíveis em tempo real?',
        'Como é feito o controle de quem pagou o ingresso na entrada do evento? (lista impressa, carimbo, outro?)',
        'Os pagamentos são feitos presencialmente, via PIX, transferência? Como é registrado que o pagamento foi recebido?',
        'Já houve situações de conflito de reserva, mesa dupla vendida ou ingresso sem controle?',
        'Como é feita a prestação de contas financeira após um evento?',
    ]),
    ('Bloco 5 — Perfil dos usuários do sistema', [
        'Quem você imagina usando o painel administrativo no dia a dia? Essa pessoa tem familiaridade com sistemas digitais?',
        'E os associados — qual faixa etária predomina? Eles usam smartphone com facilidade?',
        'Há associados que precisariam de ajuda para usar um aplicativo? Como a entidade lida com isso hoje?',
    ]),
    ('Bloco 6 — Expectativas e prioridades', [
        'Se você pudesse resolver só um problema com esse sistema, qual seria?',
        'Das funcionalidades previstas (associados, mensalidades, eventos, reservas, app do associado) — qual você considera mais urgente?',
        'Tem alguma coisa que não mencionamos mas que seria importante ter?',
        'Existe alguma preocupação sua em relação à adoção do sistema pela equipe ou pelos associados?',
        'O que você esperaria ver funcionando no sistema para considerar que foi um sucesso?',
    ]),
]

for heading, questions in blocks:
    add_block_heading(doc, heading)
    for q in questions:
        add_question_with_answer_space(doc, q)

doc.save(r'PROJETO-TCC\entrevista_cpf_pia_do_sul_respostas.docx')
print('Word gerado com sucesso')
